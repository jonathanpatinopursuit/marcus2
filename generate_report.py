import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

SRC = '/Users/jonathan/Downloads/Sample - Superstore - cleaned.csv'
OUT = '/Users/jonathan/Downloads/sales_report.xlsx'

df = pd.read_csv(SRC, parse_dates=['Order Date', 'Ship Date'])
df['Month'] = df['Order Date'].dt.to_period('M').astype(str)

# 1. Sales by Region & Category
region_category = df.groupby(['Region', 'Category'])['Sales'].sum().reset_index()

# 2. Performance by Product, Region, Segment, Time
perf = df.groupby(['Product Name', 'Region', 'Segment', 'Month']).agg(
    Total_Sales=('Sales', 'sum'), Total_Profit=('Profit', 'sum')).reset_index()

# 3. Monthly summary
monthly_summary = df.groupby('Month').agg(Sales=('Sales', 'sum'), Profit=('Profit', 'sum')).reset_index()

# 4. Underperformers (bottom 10% by profit)
underperform = df.groupby(['Category', 'Region'])['Profit'].sum().reset_index()
underperform = underperform[underperform['Profit'] < underperform['Profit'].quantile(0.1)]

# 5. Discount impact on profit
discount_impact = df.groupby('Discount %').agg(
    Avg_Profit=('Profit', 'mean'), Orders=('Order ID', 'count')).reset_index()

# 6. Region x Category pivot (wide format, complements Region_Category)
region_pivot = df.pivot_table(values='Sales', index='Region', columns='Category', aggfunc='sum', fill_value=0)

# 7. Top 10 products by sales
top_products = df.groupby('Product Name')['Sales'].sum().sort_values(ascending=False).head(10).reset_index()

# Export report
with pd.ExcelWriter(OUT, engine='openpyxl') as writer:
    region_category.to_excel(writer, sheet_name='Region_Category', index=False)
    perf.to_excel(writer, sheet_name='Product_Performance', index=False)
    monthly_summary.to_excel(writer, sheet_name='Monthly_Trend', index=False)
    underperform.to_excel(writer, sheet_name='Underperformers', index=False)
    discount_impact.to_excel(writer, sheet_name='Discount_Impact', index=False)
    region_pivot.to_excel(writer, sheet_name='Region_by_Category')
    top_products.to_excel(writer, sheet_name='Top_Products', index=False)

# Add bar chart for Top Products
wb = load_workbook(OUT)
ws = wb['Top_Products']
chart = BarChart()
chart.title = "Top 10 Products by Sales"
data = Reference(ws, min_col=2, min_row=1, max_row=11)
cats = Reference(ws, min_col=1, min_row=2, max_row=11)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
ws.add_chart(chart, "E2")
wb.save(OUT)

print(f"Report written to {OUT}")
